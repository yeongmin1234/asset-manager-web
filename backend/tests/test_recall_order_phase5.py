"""Phase 5 SCM order export and state isolation."""

import unittest
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.routers.recall_preview import router
from app.core.auth import get_current_user
from app.db.database import get_db
from app.models.menu_visibility_setting import MenuVisibilitySetting
from app.models.recall_application import (
    APPLICATION_RECEIVED, IN_PROGRESS, ORDER_CONFIRMED, ORDER_EXPORTED,
    ORDER_PENDING, RecallApplication, RecallApplicationUploadBatch,
    RecallOrderBatch, RecallStatusHistory,
)
from app.models.user import User
from app.services.recall_order_excel import SCM_HEADERS


ROOT = "/online/recall/applications/orders"


class RecallOrderPhase5Test(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        for table in (User.__table__, MenuVisibilitySetting.__table__, RecallApplicationUploadBatch.__table__,
                      RecallOrderBatch.__table__, RecallApplication.__table__, RecallStatusHistory.__table__):
            table.create(self.engine, checkfirst=True)
        self.db = Session(self.engine)
        self.db.add(User(id=1, username="operator", password_hash="test", name="담당자", role="user",
                         menu_permissions=["online_recall"]))
        self.db.add(User(id=2, username="admin", password_hash="test", name="관리자", role="admin",
                         menu_permissions=[]))
        self.db.flush()
        self.db.add(RecallApplicationUploadBatch(
            id=1, source_filename="test.xlsx", source_sha256="a" * 64,
            total_rows=4, valid_count=4, uploaded_by=1,
        ))
        self.db.flush()
        for number, status, agreement, quantity in (
            (1, IN_PROGRESS, "동의", 2),
            (2, IN_PROGRESS, "동의", 1),
            (3, APPLICATION_RECEIVED, "동의", 1),
            (4, IN_PROGRESS, "미동의", 1),
        ):
            self.db.add(RecallApplication(
                id=number, upload_batch_id=1, source_row_number=number + 1,
                application_date=date(2026, 9, 22), quantity=quantity,
                customer_name="고객{}".format(number), phone_original="01099810165",
                phone_normalized="01099810165", address="원효로 138",
                memo="방문 전 연락", replacement_shipping_agreement=agreement,
                current_status=status, created_by=1,
            ))
        self.db.commit()
        self.user = SimpleNamespace(id=1, role="user", menu_permissions=["online_recall"])
        self.app = FastAPI()
        self.app.include_router(router)
        self.app.dependency_overrides[get_current_user] = lambda: self.user
        self.app.dependency_overrides[get_db] = lambda: self.db
        self.audit_patch = patch("app.api.routers.recall_preview.record_audit_log", return_value=True)
        self.audit = self.audit_patch.start()
        self.addCleanup(self.audit_patch.stop)

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def count_batches(self):
        return self.db.scalar(select(func.count()).select_from(RecallOrderBatch))

    def test_pending_filters_preview_and_permission(self):
        with TestClient(self.app) as client:
            summary = client.get(ROOT + "/summary")
            listing = client.get(ROOT)
            preview = client.post(ROOT + "/preview", json={"ids": [1, 2]})
            invalid = client.post(ROOT + "/preview", json={"ids": [1, 3]})
            self.user.menu_permissions = []
            forbidden = client.post(ROOT + "/export", json={"ids": [1]})
        self.assertEqual(summary.json(), {"pending_count": 2, "exported_count": 0, "confirmed_count": 0})
        self.assertEqual([item["id"] for item in listing.json()["items"]], [2, 1])
        self.assertEqual(preview.status_code, 200)
        self.assertEqual((preview.json()["item_count"], preview.json()["total_quantity"]), (2, 3))
        self.assertEqual(preview.json()["columns"], list(SCM_HEADERS))
        self.assertEqual(preview.json()["rows"][0]["수취인핸드폰"], "010-9981-0165")
        self.assertEqual((invalid.status_code, forbidden.status_code), (400, 403))
        self.assertEqual(self.count_batches(), 0)
        self.assertEqual(self.db.get(RecallApplication, 1).order_status, ORDER_PENDING)

    def test_multiple_export_exact_excel_and_duplicate_prevention(self):
        with TestClient(self.app) as client:
            response = client.post(ROOT + "/export", json={"ids": [1, 2]})
            summary = client.get(ROOT + "/summary")
            duplicate = client.post(ROOT + "/export", json={"ids": [1]})
            batch_id = self.db.get(RecallApplication, 1).order_batch_id
            redownload = client.get(ROOT + "/batches/{}/download".format(batch_id))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(book.active.values)
        self.assertEqual(rows[0], SCM_HEADERS)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1], (
            "퓨어탈취필터[무상교체]", None, 2, None, "고객1", "고객1",
            "010-9981-0165", "010-9981-0165", "원효로 138", "방문 전 연락",
            None, None, "리콜", None, None, None, "무상교체",
        ))
        self.assertEqual(rows[2][2], 1)
        self.assertEqual((duplicate.status_code, self.count_batches()), (400, 1))
        self.assertEqual(redownload.status_code, 200)
        self.assertEqual(redownload.content, response.content)
        self.assertEqual(list(load_workbook(BytesIO(redownload.content), read_only=True).active.values)[0], SCM_HEADERS)
        self.assertEqual(summary.json(), {"pending_count": 0, "exported_count": 2, "confirmed_count": 0})
        self.assertEqual(self.db.get(RecallApplication, 1).current_status, IN_PROGRESS)
        self.assertEqual(self.db.get(RecallApplication, 1).order_status, ORDER_EXPORTED)
        self.assertEqual(self.db.get(RecallApplication, 2).order_batch_id, batch_id)
        batch = self.db.get(RecallOrderBatch, batch_id)
        self.assertEqual((batch.item_count, batch.total_quantity, batch.created_by), (2, 3, 1))
        audit_kwargs = self.audit.call_args.kwargs
        self.assertEqual(audit_kwargs["target_type"], "recall_order_batch")
        self.assertNotIn("고객1", str(audit_kwargs))
        self.assertNotIn("01099810165", str(audit_kwargs))

    def test_excel_failure_rolls_back_order_and_batch(self):
        with patch("app.services.recall_order_service.build_scm_workbook", side_effect=RuntimeError("broken")):
            with TestClient(self.app) as client:
                response = client.post(ROOT + "/export", json={"ids": [1, 2]})
        self.assertEqual(response.status_code, 500)
        self.assertEqual(self.count_batches(), 0)
        self.assertEqual([self.db.get(RecallApplication, item).order_status for item in (1, 2)],
                         [ORDER_PENDING, ORDER_PENDING])
        self.audit.assert_not_called()

    def test_hidden_menu_blocks_phase5_api(self):
        self.db.add(MenuVisibilitySetting(menu_key="online_recall", visible=False))
        self.db.commit()
        with TestClient(self.app) as client:
            summary = client.get(ROOT + "/summary")
            export = client.post(ROOT + "/export", json={"ids": [1]})
        self.assertEqual((summary.status_code, export.status_code), (403, 403))
        self.assertEqual(self.count_batches(), 0)

    def test_manual_admin_confirmation_only_after_export(self):
        with TestClient(self.app) as client:
            premature = client.patch(ROOT + "/1/confirm")
            self.user.role = "admin"
            pending_admin = client.patch(ROOT + "/1/confirm")
            self.user.role = "user"
            self.assertEqual(client.post(ROOT + "/export", json={"ids": [1]}).status_code, 200)
            non_admin = client.patch(ROOT + "/1/confirm")
            self.user.role = "admin"
            confirmed = client.patch(ROOT + "/1/confirm")
            repeated = client.patch(ROOT + "/1/confirm")
            summary = client.get(ROOT + "/summary")
        self.assertEqual((premature.status_code, pending_admin.status_code, non_admin.status_code,
                          confirmed.status_code, repeated.status_code), (403, 400, 403, 200, 400))
        self.assertEqual(self.db.get(RecallApplication, 1).order_status, ORDER_CONFIRMED)
        self.assertIsNotNone(self.db.get(RecallApplication, 1).order_confirmed_at)
        self.assertEqual(summary.json(), {"pending_count": 1, "exported_count": 0, "confirmed_count": 1})
        self.assertEqual(self.db.get(RecallApplication, 1).current_status, IN_PROGRESS)
        audit_kwargs = self.audit.call_args.kwargs
        self.assertEqual(audit_kwargs["target_type"], "recall_order")
        self.assertNotIn("고객1", str(audit_kwargs))


if __name__ == "__main__":
    unittest.main()
