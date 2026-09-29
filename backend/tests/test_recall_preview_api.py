import unittest
from io import BytesIO
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.routers.recall_preview import router
from app.core.auth import get_current_user
from app.db.database import get_db
from app.models.menu_visibility_setting import MenuVisibilitySetting
from app.models.recall_application import RecallApplication, RecallApplicationUploadBatch, RecallStatusHistory, RecallDuplicateResolutionHistory
from app.models.user import User
from app.services.recall_application_excel import EXCEL_COLUMNS


def sample_excel():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "자산업로드양식"
    sheet.append(list(EXCEL_COLUMNS))
    sheet.append([
        "2026.09.22", 1, "신성아", "01099810165", "원효로 138",
        "방문 전 연락부탁드립니다.", None, None, "불가(폐기", None, "동의",
    ])
    sheet.cell(3, 5, "주소 오류 체크")
    sheet.cell(4, 1)
    sheet.cell(5, 1, "1. 연락처, 시리얼번호 중복값 체크")
    sheet.cell(6, 1, "2. 주소 오류 체크")
    sheet.cell(7, 1, "3. 진행 중과 처리 완료 신규 입력건 중복 값 확인 후 카테고리 이동")
    sheet.cell(8, 11)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class RecallPreviewApiTest(unittest.TestCase):
    def test_full_application_openapi_registers_preview_once(self):
        from app.main import app

        path = "/online/recall/applications/preview"
        openapi = app.openapi()
        self.assertIn(path, openapi["paths"])
        self.assertEqual(list(openapi["paths"][path]), ["post"])
        self.assertEqual(sum(route.path == path for route in app.routes), 1)

    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        for table in (User.__table__, MenuVisibilitySetting.__table__, RecallApplicationUploadBatch.__table__, RecallApplication.__table__, RecallStatusHistory.__table__, RecallDuplicateResolutionHistory.__table__):
            table.create(self.engine, checkfirst=True)
        self.db = Session(self.engine)
        self.app = FastAPI()
        self.app.include_router(router)
        self.app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
            id=1, username="tester", name="테스터", role="user", menu_permissions=["online_recall"]
        )
        self.app.dependency_overrides[get_db] = lambda: self.db
        self.client = TestClient(self.app)

    def tearDown(self):
        self.client.close()
        self.db.close()
        self.engine.dispose()

    def upload(self, filename, content):
        return self.client.post(
            "/online/recall/applications/preview",
            files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    def test_sample_preview_includes_review_and_excluded_rows(self):
        response = self.upload("sample.xlsx", sample_excel())
        self.assertEqual(response.status_code, 200)
        result = response.json()
        self.assertEqual(result["sheet_name"], "자산업로드양식")
        self.assertEqual(result["summary"], {
            "total_rows": 7, "valid": 0, "duplicate": 0, "error": 0,
            "review": 1, "excluded": 6,
        })
        self.assertEqual(result["rows"][0]["status"], "review")
        self.assertEqual(result["rows"][0]["data"]["customer_name"], "신성아")
        self.assertEqual(result["rows"][0]["data"]["application_date"], "2026-09-22")
        self.assertEqual(result["rows"][0]["data"]["phone_normalized"], "01099810165")
        self.assertEqual(result["rows"][0]["issues"][0]["code"], "REVIEW_REQUIRED")
        self.assertEqual([row["row_kind"] for row in result["rows"][1:]],
                         ["instruction", "blank", "instruction", "instruction", "instruction", "blank"])

    def test_wrong_extension_and_unreadable_file_return_400(self):
        self.assertEqual(self.upload("sample.xls", sample_excel()).status_code, 400)
        response = self.upload("sample.xlsx", b"not a workbook")
        self.assertEqual(response.status_code, 400)
        self.assertIn(".xlsx", response.json()["detail"])

    def test_oversized_file_returns_413(self):
        response = self.upload("sample.xlsx", b"x" * (5 * 1024 * 1024 + 1))
        self.assertEqual(response.status_code, 413)

    def test_excessive_declared_sheet_range_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(EXCEL_COLUMNS))
        sheet.cell(20002, 1, "x")
        output = BytesIO()
        workbook.save(output)
        response = self.upload("huge-range.xlsx", output.getvalue())
        self.assertEqual(response.status_code, 400)

    def test_online_recall_permission_is_required(self):
        self.app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
            role="user", menu_permissions=["dashboard"]
        )
        self.assertEqual(self.upload("sample.xlsx", sample_excel()).status_code, 403)

    def test_global_visibility_off_blocks_preview_even_with_permission(self):
        self.db.add(MenuVisibilitySetting(menu_key="online_recall", visible=False))
        self.db.commit()
        self.assertEqual(self.upload("sample.xlsx", sample_excel()).status_code, 403)


if __name__ == "__main__":
    unittest.main()
