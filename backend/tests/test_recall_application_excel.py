import unittest
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel

from app.services.recall_application_excel import (
    EXCEL_COLUMNS,
    RecallApplicationExcelError,
    normalize_phone,
    preview_recall_applications,
)


HEADERS = list(EXCEL_COLUMNS)


def application_row(**changes):
    values = {
        "신청일자": "2026.09.22",
        "수량": 1,
        "성함": "홍길동",
        "*연락처": "010-9981-0165",
        "주소지": "서울시 마포구 월드컵로 10",
        "메모": "방문 전 연락",
        "*시리얼번호": "SER-001",
        "LOT 번호": "LOT-A",
        "기존 필터 회수 동의": "동의",
        "회수 일자": "2026/09/23",
        "대체 필터 출고 동의": "가능",
    }
    values.update(changes)
    return [values[header] for header in HEADERS]


def make_workbook(rows, headers=HEADERS):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "신청 데이터"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def codes(row):
    return [issue.code for issue in row.issues]


class RecallApplicationExcelTest(unittest.TestCase):
    def test_reported_a1_k8_business_sample_layout(self):
        """Recreate the supplied layout; the original workbook was not attached."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "자산업로드양식"
        sheet.append(HEADERS)
        sheet.append([
            "2026.09.22", 1, "신성아", "01099810165", "원효로 138",
            "방문 전 연락부탁드립니다.", None, None, "불가(폐기", None, "동의",
        ])
        sheet.cell(3, 5, "주소 오류 체크")
        sheet.cell(4, 1)  # completely blank row between the notes
        sheet.cell(5, 1, "1. 연락처, 시리얼번호 중복값 체크")
        sheet.cell(6, 1, "2. 주소 오류 체크")
        sheet.cell(7, 1, "3. 진행 중과 처리 완료 신규 입력건 중복 값 확인 후 카테고리 이동")
        sheet.cell(8, 11)  # retain the stated A1:K8 range
        output = BytesIO()
        workbook.save(output)

        preview = preview_recall_applications(output.getvalue(), source_filename="reported-sample.xlsx")

        self.assertEqual(preview.sheet_name, "자산업로드양식")
        self.assertEqual(preview.header_row_number, 1)
        self.assertEqual(preview.matched_columns, {field: label for label, field in EXCEL_COLUMNS.items()})
        self.assertEqual(preview.counts()["total"], 1)
        self.assertEqual(preview.counts()["review"], 1)
        self.assertEqual(preview.counts()["error"], 0)
        self.assertEqual(preview.counts()["instruction"], 4)
        self.assertEqual(preview.counts()["blank"], 2)
        self.assertEqual(
            [(row.raw_row_number, row.row_kind) for row in preview.rows if row.status == "excluded"],
            [(3, "instruction"), (4, "blank"), (5, "instruction"),
             (6, "instruction"), (7, "instruction"), (8, "blank")],
        )
        row = preview.rows[0]
        self.assertEqual(row.raw_row_number, 2)
        self.assertEqual(row.status, "review")
        self.assertEqual(codes(row), ["REVIEW_REQUIRED"])
        self.assertEqual(row.issues[0].field, "serial_number")
        self.assertEqual(row.data, {
            "application_date": date(2026, 9, 22),
            "quantity": 1,
            "customer_name": "신성아",
            "phone_original": "01099810165",
            "address": "원효로 138",
            "memo": "방문 전 연락부탁드립니다.",
            "serial_number": None,
            "lot_number": None,
            "pickup_agreement": "불가(폐기",
            "pickup_date": None,
            "replacement_shipping_agreement": "동의",
            "phone_normalized": "01099810165",
        })

    def test_valid_row_mapping_and_raw_provenance(self):
        preview = preview_recall_applications(
            make_workbook([application_row()]), source_filename="applications.xlsx",
            upload_batch="future-batch", uploaded_by="tester",
        )
        row = preview.rows[0]
        self.assertEqual(row.status, "valid")
        self.assertEqual(row.raw_row_number, 2)
        self.assertEqual(row.raw_data["*연락처"], "010-9981-0165")
        self.assertEqual(row.data["phone_original"], "010-9981-0165")
        self.assertEqual(row.data["phone_normalized"], "01099810165")
        self.assertEqual(row.data["application_date"], date(2026, 9, 22))
        self.assertEqual(row.data["quantity"], 1)
        self.assertEqual(row.data["pickup_agreement"], "동의")
        self.assertEqual(row.data["replacement_shipping_agreement"], "가능")
        self.assertEqual(preview.matched_columns["serial_number"], "*시리얼번호")
        self.assertEqual(preview.counts()["valid"], 1)
        self.assertEqual(preview.upload_batch, "future-batch")

    def test_blank_instruction_and_incomplete_data_are_distinct(self):
        empty = [None] * len(HEADERS)
        note = ["연락처, 시리얼번호 중복값 체크"] + [None] * (len(HEADERS) - 1)
        incomplete = application_row(**{"*연락처": None, "주소지": None})
        preview = preview_recall_applications(make_workbook([empty, note, incomplete]))
        self.assertEqual([row.row_kind for row in preview.rows], ["blank", "instruction", "data"])
        self.assertEqual([row.status for row in preview.rows], ["excluded", "excluded", "error"])
        self.assertEqual(preview.rows[2].raw_row_number, 4)
        self.assertEqual(preview.counts()["total"], 1)
        self.assertIn("MISSING_REQUIRED", codes(preview.rows[2]))

    def test_phone_hyphens_spaces_and_invalid_characters(self):
        self.assertEqual(normalize_phone("010 9981 0165"), "01099810165")
        self.assertEqual(normalize_phone("010-9981-0165"), "01099810165")
        preview = preview_recall_applications(make_workbook([
            application_row(**{"*연락처": "010 9981 0165", "*시리얼번호": "S1"}),
            application_row(**{"*연락처": "010-1234-5678", "*시리얼번호": "S2"}),
            application_row(**{"*연락처": "010-12ab-5678", "*시리얼번호": "S3"}),
        ]))
        self.assertEqual([row.status for row in preview.rows], ["valid", "valid", "error"])
        self.assertIn("INVALID_PHONE", codes(preview.rows[2]))
        self.assertEqual(preview.rows[2].data["phone_normalized"], "01012ab5678")

    def test_serial_duplicate_precedes_phone_duplicate_without_overwrite(self):
        preview = preview_recall_applications(make_workbook([
            application_row(**{"*시리얼번호": "SAME"}),
            application_row(**{"*시리얼번호": "same", "*연락처": "010 9981 0165"}),
        ]))
        self.assertEqual(preview.rows[1].status, "duplicate")
        self.assertEqual(codes(preview.rows[1]), ["DUPLICATE_SERIAL", "DUPLICATE_PHONE"])
        self.assertEqual(preview.rows[1].issues[0].source, "file")
        self.assertEqual(preview.rows[1].issues[0].matched_row_number, 2)
        self.assertEqual(preview.rows[0].status, "valid")

    def test_existing_duplicates_keep_status_and_source(self):
        preview = preview_recall_applications(
            make_workbook([application_row()]),
            existing_records=[{
                "id": 42, "status": "completed", "serial_number": "ser-001",
                "phone_original": "010 9981 0165",
            }],
        )
        self.assertEqual(codes(preview.rows[0]), ["DUPLICATE_SERIAL", "DUPLICATE_PHONE"])
        self.assertEqual(preview.rows[0].issues[0].source, "existing")
        self.assertEqual(preview.rows[0].issues[0].existing_id, 42)
        self.assertEqual(preview.rows[0].issues[0].existing_status, "completed")

    def test_multiple_existing_matches_keep_each_workflow_status(self):
        preview = preview_recall_applications(
            make_workbook([application_row()]),
            existing_records=[
                {"id": 1, "status": "in_progress", "serial_number": "SER-001"},
                {"id": 2, "status": "completed", "serial_number": "SER-001"},
            ],
        )
        serial_issues = [issue for issue in preview.rows[0].issues if issue.code == "DUPLICATE_SERIAL"]
        self.assertEqual([(issue.existing_id, issue.existing_status) for issue in serial_issues],
                         [(1, "in_progress"), (2, "completed")])

    def test_missing_name_address_and_optional_serial(self):
        preview = preview_recall_applications(make_workbook([
            application_row(**{"성함": None, "*시리얼번호": "S1"}),
            application_row(**{"주소지": None, "*시리얼번호": "S2", "*연락처": "01012345678"}),
            application_row(**{"*시리얼번호": None, "*연락처": "01011112222"}),
        ]))
        self.assertEqual([row.status for row in preview.rows], ["error", "error", "review"])
        self.assertEqual(preview.rows[0].issues[0].field, "customer_name")
        self.assertEqual(preview.rows[1].issues[0].field, "address")
        self.assertIn("REVIEW_REQUIRED", codes(preview.rows[2]))

    def test_quantity_errors_and_short_address_review(self):
        preview = preview_recall_applications(make_workbook([
            application_row(**{"수량": 0, "*시리얼번호": "S0"}),
            application_row(**{"수량": -1, "*시리얼번호": "S1", "*연락처": "01012345678"}),
            application_row(**{"수량": "둘", "*시리얼번호": "S2", "*연락처": "01011112222"}),
            application_row(**{"주소지": "미정", "*시리얼번호": "S3", "*연락처": "01033334444"}),
        ]))
        self.assertEqual([row.status for row in preview.rows], ["error", "error", "error", "review"])
        for row in preview.rows[:3]:
            self.assertIn("INVALID_QUANTITY", codes(row))
        self.assertEqual(preview.rows[3].issues[0].field, "address")

    def test_date_formats_excel_cells_and_invalid_date_preserve_raw(self):
        preview = preview_recall_applications(make_workbook([
            application_row(**{"신청일자": "2026.09.22", "*시리얼번호": "S1"}),
            application_row(**{"신청일자": "2026-09-22", "*시리얼번호": "S2", "*연락처": "01012345678"}),
            application_row(**{"신청일자": "2026/09/22", "*시리얼번호": "S3", "*연락처": "01011112222"}),
            application_row(**{"신청일자": datetime(2026, 9, 22), "*시리얼번호": "S4", "*연락처": "01033334444"}),
            application_row(**{"신청일자": to_excel(datetime(2026, 9, 22)), "*시리얼번호": "S5", "*연락처": "01055556666"}),
            application_row(**{"신청일자": "2026-99-22", "*시리얼번호": "S6", "*연락처": "01077778888"}),
        ]))
        self.assertEqual([row.data["application_date"] for row in preview.rows[:5]], [date(2026, 9, 22)] * 5)
        self.assertEqual(preview.rows[5].status, "error")
        self.assertIn("INVALID_DATE", codes(preview.rows[5]))
        self.assertEqual(preview.rows[5].raw_data["신청일자"], "2026-99-22")

    def test_missing_header_and_bad_file_fail_closed(self):
        with self.assertRaises(RecallApplicationExcelError):
            preview_recall_applications(b"not-xlsx")
        with self.assertRaisesRegex(RecallApplicationExcelError, "컬럼"):
            preview_recall_applications(make_workbook([application_row()], headers=HEADERS[:-1]))

    def test_header_after_note_and_extra_column_preserve_source_position(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["주소 오류 체크"])
        sheet.append(HEADERS + ["추가 메모"])
        sheet.append(application_row() + ["원본 추가값"])
        output = BytesIO()
        workbook.save(output)
        preview = preview_recall_applications(output.getvalue())
        self.assertEqual(preview.header_row_number, 2)
        self.assertEqual(preview.rows[0].raw_row_number, 3)
        self.assertEqual(preview.rows[0].raw_data["추가 메모"], "원본 추가값")


if __name__ == "__main__":
    unittest.main()
