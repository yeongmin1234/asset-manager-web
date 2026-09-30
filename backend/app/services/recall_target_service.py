"""Parse and register cumulative raw recall targets separately from applications."""

import math
import re
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Dict, List

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func, or_, select, text
from sqlalchemy.orm import Session

from app.models.recall_application import RecallApplication
from app.models.recall_target import RecallTarget, RecallTargetUploadBatch
from app.services.recall_target_matching import run_auto_matching


HEADERS = {
    "sales_channel": ("판매채널", "판매 채널", "채널", "주문몰"),
    "original_order_no": ("주문번호", "원주문번호", "주문 번호"),
    "customer_name": ("고객명", "성함", "이름", "수취인"),
    "phone_raw": ("연락처", "전화번호", "휴대폰번호", "수취인핸드폰"),
    "address": ("주소", "주소지", "총주소", "배송주소"),
    "delivery_message": ("배송메시지", "배송 메시지", "메시지"),
    "serial_number": ("시리얼번호", "시리얼 번호", "시리얼", "serialnumber"),
    "lot_number": ("LOT 번호", "LOT번호", "로트번호"),
    "purchase_date": ("구매일", "구매일자", "판매일", "판매일자"),
}
REQUIRED = ("customer_name", "phone_raw", "address")


def _text(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _key(value):
    return "".join((_text(value) or "").replace("*", "").split()).casefold()


def _date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch=epoch)
            return parsed.date() if isinstance(parsed, datetime) else parsed if isinstance(parsed, date) else None
        except (ValueError, OverflowError):
            return None
    match = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", _text(value) or "")
    if match:
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            pass
    return None


def _existing(db):
    records = db.scalars(select(RecallTarget).where(RecallTarget.is_deleted.is_(False)).order_by(RecallTarget.id)).all()
    indexes = {key: {} for key in ("serial_number", "phone_normalized", "original_order_no")}
    for item in records:
        for field, index in indexes.items():
            value = getattr(item, field)
            if value:
                index.setdefault(value.casefold(), (item.id, None))
    return indexes


def preview_targets(db: Session, file_bytes: bytes) -> Dict[str, object]:
    if not file_bytes:
        raise ValueError("빈 Excel 파일입니다.")
    try:
        book = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("읽을 수 없는 .xlsx 파일입니다.") from exc
    try:
        found = None
        for sheet in book.worksheets:
            for row_number, cells in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 30), values_only=True), start=1):
                normalized = [_key(value) for value in cells]
                indexes = {}
                for field, aliases in HEADERS.items():
                    matches = [i for i, value in enumerate(normalized) if value in {_key(alias) for alias in aliases}]
                    if len(matches) > 1:
                        raise ValueError("중복된 Excel 컬럼입니다: {}".format(field))
                    if matches:
                        indexes[field] = matches[0]
                if all(field in indexes for field in REQUIRED):
                    found = (sheet, row_number, indexes)
                    break
            if found:
                break
        if found is None:
            raise ValueError("고객명·연락처·주소 컬럼을 가진 시트를 찾지 못했습니다.")
        sheet, header_row, indexes = found
        if sheet.max_row - header_row > 20000 or sheet.max_column > 100:
            raise ValueError("Excel 미리보기 범위가 너무 큽니다.")
        seen = _existing(db)
        rows = []
        counts = {key: 0 for key in ("valid", "duplicate", "review", "excluded")}
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            data = {field: _text(cells[index]) if index < len(cells) else None for field, index in indexes.items()}
            data = {field: data.get(field) for field in HEADERS}
            data["phone_normalized"] = re.sub(r"\D", "", data["phone_raw"] or "") or None
            raw_date = cells[indexes["purchase_date"]] if "purchase_date" in indexes and indexes["purchase_date"] < len(cells) else None
            data["purchase_date"] = _date(raw_date, book.epoch)
            reasons = []
            limits = {"sales_channel": 100, "original_order_no": 100, "customer_name": 100,
                      "phone_raw": 100, "phone_normalized": 30, "serial_number": 100, "lot_number": 100}
            if any(data[field] and len(data[field]) > limit for field, limit in limits.items()):
                status = "excluded"
                reasons.append("VALUE_TOO_LONG")
            elif not any(_text(value) for value in cells) or not any(data[field] for field in REQUIRED):
                status = "excluded"
                reasons.append("EMPTY_OR_UNUSABLE")
            else:
                for field in REQUIRED:
                    if not data[field]:
                        reasons.append("{}_MISSING".format(field.upper()))
                if data["phone_raw"] and len(data["phone_normalized"] or "") not in (10, 11):
                    reasons.append("PHONE_INVALID")
                if data["address"] and (len(data["address"]) < 5 or data["address"].casefold() in ("미정", "없음", "주소불명", "?")):
                    reasons.append("ADDRESS_CHECK")
                if "serial_number" in indexes and not data["serial_number"]:
                    reasons.append("SERIAL_CHECK")
                if data["purchase_date"] is None and _text(raw_date):
                    reasons.append("DATE_INVALID")
                matches = []
                reference = None
                for field, code in (("serial_number", "SERIAL"), ("phone_normalized", "PHONE"), ("original_order_no", "ORDER_NO")):
                    value = data[field]
                    if value and value.casefold() in seen[field]:
                        matches.append(code)
                        if reference is None:
                            reference = seen[field][value.casefold()]
                status = "duplicate" if matches else "review" if reasons else "valid"
                for field in seen:
                    value = data[field]
                    if value:
                        seen[field].setdefault(value.casefold(), (None, row_number))
            counts[status] += 1
            rows.append({"raw_row_number": row_number, "status": status, "data": data,
                         "duplicate_reason": (matches[0] if len(matches) == 1 else "PHONE_AND_SERIAL" if set(matches) == {"PHONE", "SERIAL"} else "MULTIPLE") if status == "duplicate" else None,
                         "duplicate_reference_id": reference[0] if status == "duplicate" else None,
                         "matched_row_number": reference[1] if status == "duplicate" else None,
                         "review_required": bool(reasons) and status != "excluded", "review_reason": ",".join(reasons) if status != "excluded" else None,
                         "reasons": reasons})
        return {"sheet_name": sheet.title, "summary": {"total_rows": len(rows), **counts}, "rows": rows}
    finally:
        book.close()


def commit_targets(db: Session, file_bytes: bytes, filename: str, user_id: int) -> Dict[str, object]:
    try:
        preview = preview_targets(db, file_bytes)
        counts = preview["summary"]
        batch = RecallTargetUploadBatch(original_filename=filename[:255], total_count=counts["total_rows"],
                                        normal_count=counts["valid"], duplicate_count=counts["duplicate"],
                                        review_count=counts["review"], excluded_count=counts["excluded"], created_by=user_id)
        db.add(batch)
        db.flush()
        row_ids = {}
        for row in preview["rows"]:
            if row["status"] == "excluded":
                continue
            data = row["data"]
            reference_id = row["duplicate_reference_id"] or row_ids.get(row["matched_row_number"])
            target = RecallTarget(batch_id=batch.id, source_row_number=row["raw_row_number"], created_by=user_id,
                                  duplicate_flag=row["status"] == "duplicate", duplicate_reason=row["duplicate_reason"],
                                  duplicate_reference_id=reference_id, review_required=row["review_required"],
                                  review_reason=row["review_reason"], **data)
            db.add(target)
            db.flush()
            row_ids[row["raw_row_number"]] = target.id
        matching = run_auto_matching(db, user_id=user_id, target_ids=list(row_ids.values())) if row_ids else {"matched": 0, "review": 0, "unmatched": 0}
        db.commit()
        return {"batch_id": batch.id, "registered": len(row_ids), "normal": counts["valid"],
                "duplicate": counts["duplicate"], "review": counts["review"], "excluded": counts["excluded"],
                "total": counts["total_rows"], "matching": matching}
    except Exception:
        db.rollback()
        raise


def target_summary(db: Session):
    rows = db.execute(select(RecallTarget.duplicate_flag, RecallTarget.review_required, func.count()).where(
        RecallTarget.is_deleted.is_(False)).group_by(RecallTarget.duplicate_flag, RecallTarget.review_required)).all()
    total = sum(count for _, _, count in rows)
    return {"total_count": total, "normal_count": sum(count for duplicate, review, count in rows if not duplicate and not review),
            "duplicate_count": sum(count for duplicate, _, count in rows if duplicate),
            "review_count": sum(count for _, review, count in rows if review)}


def list_targets(db: Session, keyword: str, status: str, page: int, page_size: int,
                 match_status: str = "", sales_channel: str = ""):
    if status not in ("", "valid", "duplicate", "review"):
        raise ValueError("대상 상태 필터가 올바르지 않습니다.")
    if match_status not in ("", "MATCHED", "UNMATCHED", "REVIEW"):
        raise ValueError("신청 상태 필터가 올바르지 않습니다.")
    query = select(RecallTarget).where(RecallTarget.is_deleted.is_(False))
    if sales_channel:
        query = query.where(RecallTarget.sales_channel == sales_channel)
    if keyword:
        pattern = "%{}%".format(keyword.replace("%", "\\%").replace("_", "\\_"))
        query = query.where(or_(RecallTarget.customer_name.ilike(pattern, escape="\\"),
                                RecallTarget.phone_raw.ilike(pattern, escape="\\"),
                                RecallTarget.phone_normalized.ilike(pattern, escape="\\"),
                                RecallTarget.original_order_no.ilike(pattern, escape="\\"),
                                RecallTarget.serial_number.ilike(pattern, escape="\\")))
    if status == "valid":
        query = query.where(RecallTarget.duplicate_flag.is_(False), RecallTarget.review_required.is_(False))
    elif status == "duplicate":
        query = query.where(RecallTarget.duplicate_flag.is_(True))
    elif status == "review":
        query = query.where(RecallTarget.review_required.is_(True))
    if match_status:
        query = query.where(RecallTarget.match_status == match_status)
    total = int(db.scalar(select(func.count()).select_from(query.subquery())) or 0)
    items = db.scalars(query.order_by(RecallTarget.id.desc()).offset((page - 1) * page_size).limit(page_size)).all()
    application_ids = [item.matched_application_id for item in items if item.matched_application_id]
    applications = {item.id: item for item in db.scalars(select(RecallApplication).where(RecallApplication.id.in_(application_ids))).all()} if application_ids else {}
    results = []
    for item in items:
        data = {column.name: getattr(item, column.name) for column in RecallTarget.__table__.columns}
        application = applications.get(item.matched_application_id)
        data["application_date"] = application.application_date if application else None
        data["application_status"] = application.current_status if application else None
        results.append(data)
    return {"items": results,
            "total": total, "page": page, "page_size": page_size, "total_pages": max(1, math.ceil(total / page_size))}


def list_target_batches(db: Session):
    batches = db.scalars(select(RecallTargetUploadBatch).order_by(RecallTargetUploadBatch.id.desc())).all()
    return [{column.name: getattr(batch, column.name) for column in RecallTargetUploadBatch.__table__.columns} for batch in batches]


def bulk_soft_delete_targets(db: Session, *, ids, reason: str, user_id: int):
    """Soft delete selected targets while retaining application and audit history."""
    if not ids or len(ids) > 100 or any(type(value) is not int or value < 1 for value in ids) or len(set(ids)) != len(ids):
        raise ValueError("삭제할 리콜 대상을 1~100건 중복 없이 선택해주세요.")
    cleaned_reason = (reason or "").strip()
    if not cleaned_reason or len(cleaned_reason) > 255:
        raise ValueError("삭제 사유는 1~255자로 입력해주세요.")
    try:
        if db.get_bind().dialect.name == "postgresql":
            db.execute(text("SELECT pg_advisory_xact_lock(714257001)"))
        targets = list(db.scalars(select(RecallTarget).where(
            RecallTarget.id.in_(ids), RecallTarget.is_deleted.is_(False)
        ).order_by(RecallTarget.id).with_for_update()).all())
        if len(targets) != len(ids):
            raise ValueError("선택한 대상 중 이미 삭제되었거나 찾을 수 없는 건이 있습니다.")
        deleted_at = datetime.now(timezone.utc)
        for target in targets:
            target.is_deleted = True
            target.deleted_at = deleted_at
            target.deleted_by = user_id
            target.delete_reason = cleaned_reason
        db.commit()
        return {"deleted": len(targets), "ids": list(ids), "deleted_at": deleted_at,
                "deleted_by": user_id, "delete_reason": cleaned_reason}
    except Exception:
        db.rollback()
        raise
