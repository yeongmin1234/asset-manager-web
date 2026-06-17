from datetime import date, datetime, timezone
from io import BytesIO
import re
from typing import Dict, List, Optional, Set, Union

from sqlalchemy import or_, select
from sqlalchemy.orm import Session, joinedload
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.asset import Asset, AssetStatus
from app.models.category import Category
from app.models.department import Department
from app.models.history import AssetActionType
from app.schemas.asset import (
    AssetCreate,
    AssetImportCommitResponse,
    AssetImportData,
    AssetImportPreviewResponse,
    AssetImportPreviewRow,
    AssetUpdate,
)
from app.services.history_service import (
    record_asset_created,
    record_asset_deleted,
    record_asset_disposed,
    record_asset_history,
)
from app.services.activity_log_service import (
    record_asset_activity,
    serialize_asset_activity_data,
)


class AssetNotFoundError(Exception):
    pass


class AssetConflictError(Exception):
    pass


class AssetValidationError(Exception):
    pass


IMPORT_HEADERS = [
    "제품명",
    "분류",
    "부서(사용자명)",
    "상태",
    "시리얼번호",
    "메모",
    "구매일",
    "모델명",
]

ALLOWED_IMPORT_STATUSES = {status.value for status in AssetStatus}


def get_assets(
    db: Session,
    *,
    status: Optional[AssetStatus] = None,
    category_id: Optional[int] = None,
    department_id: Optional[int] = None,
    department_name: Optional[str] = None,
    keyword: Optional[str] = None,
) -> List[Asset]:
    statement = (
        select(Asset)
        .options(joinedload(Asset.category), joinedload(Asset.department))
        .outerjoin(Department, Asset.department_id == Department.id)
        .where(Asset.deleted_at.is_(None))
    )

    if status is not None:
        statement = statement.where(Asset.status == status)
    if category_id is not None:
        statement = statement.where(Asset.category_id == category_id)
    if department_id is not None:
        statement = statement.where(Asset.department_id == department_id)
    if department_name:
        department_pattern = f"%{department_name.strip()}%"
        if department_pattern != "%%":
            statement = statement.where(
                or_(
                    Asset.department_name.ilike(department_pattern),
                    Asset.user_name.ilike(department_pattern),
                    Department.name.ilike(department_pattern),
                )
            )

    if keyword:
        keyword_pattern = f"%{keyword.strip()}%"
        if keyword_pattern != "%%":
            statement = statement.where(
                or_(
                    Asset.name.ilike(keyword_pattern),
                    Asset.model_name.ilike(keyword_pattern),
                    Asset.serial_number.ilike(keyword_pattern),
                    Asset.department_name.ilike(keyword_pattern),
                    Asset.user_name.ilike(keyword_pattern),
                    Department.name.ilike(keyword_pattern),
                )
            )

    statement = statement.order_by(Asset.created_at.desc(), Asset.id.desc())
    return list(db.scalars(statement).all())


def build_assets_excel(assets: List[Asset]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "자산목록"

    headers = [
        "제품명",
        "상태",
        "분류",
        "시리얼번호",
        "부서(사용자명)",
        "구매일",
        "모델명",
        "메모",
        "등록일",
    ]
    worksheet.append(headers)

    for asset in assets:
        worksheet.append(
            [
                asset.name or "",
                format_status(asset.status),
                asset.category.name if asset.category else "",
                asset.serial_number or "",
                format_department_user(asset),
                format_excel_date(asset.purchase_date),
                asset.model_name or "",
                asset.note or "",
                format_excel_date(asset.created_at),
            ]
        )

    header_fill = PatternFill(fill_type="solid", fgColor="EAF2FF")
    header_font = Font(bold=True, color="172033")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    column_widths = [28, 12, 16, 20, 22, 14, 22, 36, 20]
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_asset_import_template() -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "자산업로드양식"
    worksheet.append(IMPORT_HEADERS)
    worksheet.append(["테스트 노트북", "노트북", "총무팀 홍길동", "사용중", "TEST001", "예시 행", "2026-06-15", "ThinkPad"])

    header_fill = PatternFill(fill_type="solid", fgColor="EAF2FF")
    header_font = Font(bold=True, color="172033")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    column_widths = [26, 16, 24, 12, 20, 32, 14, 22]
    for index, width in enumerate(column_widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def preview_assets_import(db: Session, file_bytes: bytes) -> AssetImportPreviewResponse:
    raw_rows = parse_asset_import_excel(file_bytes)
    categories_by_name = {
        category.name.strip(): category
        for category in db.scalars(select(Category).where(Category.is_active.is_(True))).all()
    }
    existing_serials = {
        serial.strip().upper()
        for serial in db.scalars(
            select(Asset.serial_number).where(
                Asset.serial_number.is_not(None),
                Asset.deleted_at.is_(None),
            )
        ).all()
        if serial
    }

    seen_serials: Set[str] = set()
    rows: List[AssetImportPreviewRow] = []

    for raw_row in raw_rows:
        row_number = raw_row["row_number"]
        errors: List[str] = []
        name = normalize_text(raw_row.get("제품명"))
        category_name = normalize_text(raw_row.get("분류"))
        department_name = normalize_text(raw_row.get("부서(사용자명)"))
        status_value = normalize_text(raw_row.get("상태"))
        serial_number = normalize_serial_number_for_import(raw_row.get("시리얼번호"), errors)
        note = normalize_text(raw_row.get("메모"))
        purchase_date = parse_import_date(raw_row.get("구매일"), errors)
        model_name = normalize_text(raw_row.get("모델명"))

        if not name:
            errors.append("제품명은 필수입니다.")
        if not category_name:
            errors.append("분류는 필수입니다.")
        elif category_name not in categories_by_name:
            errors.append(f"등록된 분류가 아닙니다: {category_name}")
        if not status_value:
            errors.append("상태는 필수입니다.")
        elif status_value not in ALLOWED_IMPORT_STATUSES:
            errors.append("상태는 사용중, 미사용, 폐기 중 하나여야 합니다.")

        if serial_number:
            if serial_number in existing_serials:
                errors.append(f"이미 등록된 시리얼번호입니다: {serial_number}")
            if serial_number in seen_serials:
                errors.append(f"엑셀 파일 안에서 중복된 시리얼번호입니다: {serial_number}")
            seen_serials.add(serial_number)

        data = AssetImportData(
            name=name,
            category_name=category_name,
            department_name=department_name,
            status=AssetStatus(status_value) if status_value in ALLOWED_IMPORT_STATUSES else status_value,
            serial_number=serial_number,
            note=note,
            purchase_date=purchase_date,
            model_name=model_name,
        )

        rows.append(
            AssetImportPreviewRow(
                row_number=row_number,
                is_valid=not errors,
                data=data,
                errors=errors,
            )
        )

    valid_rows = sum(1 for row in rows if row.is_valid)
    return AssetImportPreviewResponse(
        total_rows=len(rows),
        valid_rows=valid_rows,
        error_rows=len(rows) - valid_rows,
        rows=rows,
    )


def commit_assets_import(
    db: Session,
    rows: List[AssetImportPreviewRow],
) -> AssetImportCommitResponse:
    errors: List[AssetImportPreviewRow] = []
    created_count = 0
    seen_serials: Set[str] = set()

    for row in rows:
        data = row.data
        if not row.is_valid or data is None:
            errors.append(
                AssetImportPreviewRow(
                    row_number=row.row_number,
                    is_valid=False,
                    data=data,
                    errors=row.errors or ["미리보기에서 오류가 있는 행은 등록할 수 없습니다."],
                )
            )
            continue

        row_errors = validate_import_commit_row(db, data, seen_serials)
        if row_errors:
            errors.append(
                AssetImportPreviewRow(
                    row_number=row.row_number,
                    is_valid=False,
                    data=data,
                    errors=row_errors,
                )
            )
            continue

        category = db.scalar(
            select(Category).where(
                Category.name == data.category_name,
                Category.is_active.is_(True),
            )
        )

        try:
            create_asset(
                db,
                AssetCreate(
                    category_id=category.id,
                    department_id=None,
                    department_name=data.department_name,
                    name=data.name,
                    model_name=data.model_name,
                    serial_number=data.serial_number,
                    purchase_date=data.purchase_date,
                    purchase_price=None,
                    user_name=None,
                    status=AssetStatus(data.status),
                    note=data.note,
                ),
            )
            created_count += 1
        except (AssetConflictError, AssetValidationError) as exc:
            db.rollback()
            errors.append(
                AssetImportPreviewRow(
                    row_number=row.row_number,
                    is_valid=False,
                    data=data,
                    errors=[str(exc)],
                )
            )

    return AssetImportCommitResponse(
        created_count=created_count,
        skipped_count=len(rows) - created_count,
        errors=errors,
    )


def validate_import_commit_row(
    db: Session,
    data: AssetImportData,
    seen_serials: Set[str],
) -> List[str]:
    errors: List[str] = []
    status_value = data.status.value if isinstance(data.status, AssetStatus) else data.status

    if not data.name:
        errors.append("제품명은 필수입니다.")
    if not data.category_name:
        errors.append("분류는 필수입니다.")
    elif db.scalar(
        select(Category.id).where(
            Category.name == data.category_name,
            Category.is_active.is_(True),
        )
    ) is None:
        errors.append(f"등록된 분류가 아닙니다: {data.category_name}")
    if not status_value:
        errors.append("상태는 필수입니다.")
    elif status_value not in ALLOWED_IMPORT_STATUSES:
        errors.append("상태는 사용중, 미사용, 폐기 중 하나여야 합니다.")

    if data.serial_number:
        serial_number = data.serial_number.strip().upper()
        data.serial_number = serial_number
        if re.fullmatch(r"[A-Z0-9]+", serial_number) is None:
            errors.append("시리얼번호는 영문과 숫자만 입력할 수 있습니다.")
        elif serial_number in seen_serials:
            errors.append(f"엑셀 파일 안에서 중복된 시리얼번호입니다: {serial_number}")
        else:
            seen_serials.add(serial_number)
            try:
                ensure_serial_number_is_available(db, serial_number)
            except AssetConflictError:
                errors.append(f"이미 등록된 시리얼번호입니다: {serial_number}")

    return errors


def parse_asset_import_excel(file_bytes: bytes) -> List[Dict[str, object]]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=True)
    worksheet = workbook.active
    header_values = [normalize_text(cell.value) for cell in worksheet[1]]
    header_indexes = {header: index for index, header in enumerate(header_values)}
    missing_headers = [header for header in IMPORT_HEADERS if header not in header_indexes]
    if missing_headers:
        raise AssetValidationError(f"엑셀 양식 컬럼이 없습니다: {', '.join(missing_headers)}")

    rows: List[Dict[str, object]] = []
    for row_index in range(2, worksheet.max_row + 1):
        row_values = [worksheet.cell(row=row_index, column=column).value for column in range(1, worksheet.max_column + 1)]
        if all(value is None or str(value).strip() == "" for value in row_values):
            continue

        rows.append(
            {
                "row_number": row_index,
                **{
                    header: row_values[header_indexes[header]]
                    for header in IMPORT_HEADERS
                },
            }
        )
    return rows


def normalize_text(value: object) -> Optional[str]:
    if value is None:
        return None
    normalized_value = str(value).strip()
    return normalized_value or None


def normalize_serial_number_for_import(value: object, errors: List[str]) -> Optional[str]:
    serial_number = normalize_text(value)
    if not serial_number:
        return None

    normalized_serial_number = serial_number.upper()
    if re.fullmatch(r"[A-Z0-9]+", normalized_serial_number) is None:
        errors.append("시리얼번호는 영문과 숫자만 입력할 수 있습니다.")
        return normalized_serial_number
    return normalized_serial_number


def parse_import_date(value: object, errors: List[str]) -> Optional[date]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    value_text = str(value).strip()
    try:
        return datetime.strptime(value_text, "%Y-%m-%d").date()
    except ValueError:
        errors.append("구매일은 YYYY-MM-DD 형식이어야 합니다.")
        return None


def format_status(status: Optional[Union[AssetStatus, str]]) -> str:
    if status is None:
        return ""
    return status.value if isinstance(status, AssetStatus) else str(status)


def format_department_user(asset: Asset) -> str:
    if asset.department_name:
        return asset.department_name
    if asset.department and asset.department.name:
        return asset.department.name
    return asset.user_name or ""


def format_excel_date(value: Optional[Union[date, datetime]]) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None).strftime("%Y-%m-%d %H:%M:%S")
    return value.isoformat()


def get_asset(db: Session, asset_id: int) -> Asset:
    asset = db.scalar(
        select(Asset).where(
            Asset.id == asset_id,
            Asset.deleted_at.is_(None),
        )
    )
    if asset is None:
        raise AssetNotFoundError(f"Asset not found: {asset_id}")
    return asset


def create_asset(
    db: Session,
    asset_create: AssetCreate,
    *,
    actor_ip: Optional[str] = None,
    user_agent: Optional[str] = None,
) -> Asset:
    validate_category(db, asset_create.category_id)
    if asset_create.department_id is not None:
        validate_department(db, asset_create.department_id)
    if asset_create.serial_number is not None:
        ensure_serial_number_is_available(db, asset_create.serial_number)

    asset = Asset(**asset_create.model_dump())

    try:
        db.add(asset)
        db.flush()
        record_asset_created(db, asset.id)
        record_asset_activity(
            db,
            action_type="등록",
            target_id=asset.id,
            target_name=asset.name,
            actor_ip=actor_ip,
            user_agent=user_agent,
            summary=f"자산 등록: {asset.name}",
            after_data=serialize_asset_activity_data(asset),
        )
        db.commit()
        db.refresh(asset)
    except Exception:
        db.rollback()
        raise

    return asset


def update_asset(
    db: Session,
    asset_id: int,
    asset_update: AssetUpdate,
    *,
    actor_ip: Optional[str] = None,
    user_agent: Optional[str] = None,
) -> Asset:
    asset = get_asset(db, asset_id)
    update_data = asset_update.model_dump(exclude_unset=True)
    before_data = serialize_asset_activity_data(asset)

    if "category_id" in update_data:
        validate_category(db, update_data["category_id"])
    if update_data.get("department_id") is not None:
        validate_department(db, update_data["department_id"])
    if "serial_number" in update_data and update_data["serial_number"] is not None:
        ensure_serial_number_is_available(
            db,
            update_data["serial_number"],
            exclude_asset_id=asset.id,
        )

    try:
        for field_name, new_value in update_data.items():
            old_value = getattr(asset, field_name)
            if old_value == new_value:
                continue

            setattr(asset, field_name, new_value)
            action_type = (
                AssetActionType.STATUS_CHANGED
                if field_name == "status"
                else AssetActionType.UPDATED
            )
            record_asset_history(
                db,
                asset_id=asset.id,
                action_type=action_type,
                field_name=field_name,
                old_value=old_value,
                new_value=new_value,
            )

        db.flush()
        record_asset_activity(
            db,
            action_type="수정",
            target_id=asset.id,
            target_name=asset.name,
            actor_ip=actor_ip,
            user_agent=user_agent,
            summary=f"자산 수정: {asset.name}",
            before_data=before_data,
            after_data=serialize_asset_activity_data(asset),
        )
        db.commit()
        db.refresh(asset)
    except Exception:
        db.rollback()
        raise

    return asset


def dispose_asset(
    db: Session,
    asset_id: int,
    *,
    actor_ip: Optional[str] = None,
    user_agent: Optional[str] = None,
) -> Asset:
    asset = get_asset(db, asset_id)
    if asset.status == AssetStatus.DISPOSED:
        return asset

    try:
        before_data = serialize_asset_activity_data(asset)
        old_status = asset.status
        asset.status = AssetStatus.DISPOSED
        record_asset_disposed(
            db,
            asset_id=asset.id,
            old_value=old_status,
            new_value=AssetStatus.DISPOSED,
        )
        db.flush()
        record_asset_activity(
            db,
            action_type="폐기",
            target_id=asset.id,
            target_name=asset.name,
            actor_ip=actor_ip,
            user_agent=user_agent,
            summary=f"자산 폐기: {asset.name}",
            before_data=before_data,
            after_data=serialize_asset_activity_data(asset),
        )
        db.commit()
        db.refresh(asset)
    except Exception:
        db.rollback()
        raise

    return asset


def soft_delete_asset(
    db: Session,
    asset_id: int,
    *,
    actor_ip: Optional[str] = None,
    user_agent: Optional[str] = None,
) -> Asset:
    asset = get_asset(db, asset_id)

    try:
        before_data = serialize_asset_activity_data(asset)
        deleted_at = datetime.now(timezone.utc)
        asset.deleted_at = deleted_at
        record_asset_deleted(db, asset.id, new_value=deleted_at)
        record_asset_activity(
            db,
            action_type="삭제",
            target_id=asset.id,
            target_name=asset.name,
            actor_ip=actor_ip,
            user_agent=user_agent,
            summary=f"자산 삭제: {asset.name}",
            before_data=before_data,
        )
        db.commit()
        db.refresh(asset)
    except Exception:
        db.rollback()
        raise

    return asset


def validate_category(db: Session, category_id: int) -> None:
    category = db.scalar(select(Category).where(Category.id == category_id))
    if category is None:
        raise AssetValidationError(f"Category not found: {category_id}")


def validate_department(db: Session, department_id: int) -> None:
    department = db.scalar(select(Department).where(Department.id == department_id))
    if department is None:
        raise AssetValidationError(f"Department not found: {department_id}")


def ensure_serial_number_is_available(
    db: Session,
    serial_number: str,
    *,
    exclude_asset_id: Optional[int] = None,
) -> None:
    statement = select(Asset.id).where(
        Asset.serial_number == serial_number,
        Asset.deleted_at.is_(None),
    )
    if exclude_asset_id is not None:
        statement = statement.where(Asset.id != exclude_asset_id)

    existing_asset = db.scalar(statement)
    if existing_asset is not None:
        raise AssetConflictError(f"serial_number already exists: {serial_number}")
