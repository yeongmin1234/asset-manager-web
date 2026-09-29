"""SCM workbook layout for recall replacement-filter orders."""

import re
from io import BytesIO
from typing import Dict, List

from openpyxl import Workbook

from app.models.recall_application import RecallApplication


SCM_HEADERS = (
    "상품명", "옵션", "주문수량", "상품코드", "수취인", "주문자", "수취인전화",
    "수취인핸드폰", "총주소", "메시지", "우편번호", "발신인", "주문몰",
    "매출액", "판매가", "주문번호", "메타태그",
)


def format_scm_phone(value: str) -> str:
    original = value or ""
    digits = re.sub(r"\D", "", original)
    if len(digits) == 11:
        return "{}-{}-{}".format(digits[:3], digits[3:7], digits[7:])
    return original


def scm_row(application: RecallApplication) -> Dict[str, object]:
    phone = format_scm_phone(application.phone_original)
    return dict(zip(SCM_HEADERS, (
        "퓨어탈취필터[무상교체]", None, application.quantity, None,
        application.customer_name, application.customer_name, phone, phone,
        application.address, application.memo or None, None, None, "리콜",
        None, None, None, "무상교체",
    )))


def build_scm_workbook(rows: List[Dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SCM 발주"
    sheet.append(SCM_HEADERS)
    for row in rows:
        sheet.append([row[header] for header in SCM_HEADERS])
        row_number = sheet.max_row
        for column, header in enumerate(SCM_HEADERS, start=1):
            value = row[header]
            if isinstance(value, str):
                # Customer-supplied text must remain text, even when it starts with '='.
                sheet.cell(row=row_number, column=column).data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
