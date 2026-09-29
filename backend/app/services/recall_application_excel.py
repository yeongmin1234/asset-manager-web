"""Read-only preview rules for the recall application/order worksheet.

This is intentionally separate from a future recall-target raw-data importer.
It does not query or write a database, and it does not decide import policy.
"""

import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable, Mapping, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


class RecallApplicationExcelError(ValueError):
    pass


EXCEL_COLUMNS = {
    "신청일자": "application_date",
    "수량": "quantity",
    "성함": "customer_name",
    "*연락처": "phone_original",
    "주소지": "address",
    "메모": "memo",
    "*시리얼번호": "serial_number",
    "LOT 번호": "lot_number",
    "기존 필터 회수 동의": "pickup_agreement",
    "회수 일자": "pickup_date",
    "대체 필터 출고 동의": "replacement_shipping_agreement",
}
REQUIRED_FIELDS = ("customer_name", "phone_original", "address")
DATE_FIELDS = ("application_date", "pickup_date")
AGREEMENT_FIELDS = ("pickup_agreement", "replacement_shipping_agreement")
MAX_PREVIEW_ROWS = 20000
MAX_PREVIEW_COLUMNS = 100
INSTRUCTION_MARKERS = (
    "중복값 체크", "중복 값 확인", "주소 오류 체크", "카테고리 이동", "진행 중과 처리 완료",
)


@dataclass(frozen=True)
class ValidationIssue:
    code: str
    field: Optional[str]
    message: str
    source: Optional[str] = None  # file or existing
    matched_row_number: Optional[int] = None
    existing_id: Optional[Any] = None
    existing_status: Optional[str] = None


@dataclass(frozen=True)
class PreviewRow:
    raw_row_number: int
    row_kind: str  # data, instruction, blank
    status: str  # valid, duplicate, error, review, excluded
    raw_data: Mapping[str, Any]
    data: Mapping[str, Any]
    issues: Tuple[ValidationIssue, ...]


@dataclass(frozen=True)
class ApplicationPreview:
    source_filename: Optional[str]
    sheet_name: str
    header_row_number: int
    matched_columns: Mapping[str, str]
    rows: Tuple[PreviewRow, ...]
    upload_batch: Optional[str] = None
    uploaded_by: Optional[str] = None
    uploaded_at: Optional[datetime] = None

    def counts(self) -> Mapping[str, int]:
        result = {key: 0 for key in ("valid", "duplicate", "error", "review", "instruction", "blank")}
        for row in self.rows:
            result[row.row_kind if row.status == "excluded" else row.status] += 1
        result["total"] = sum(result[key] for key in ("valid", "duplicate", "error", "review"))
        return result


def normalize_phone(value: Any) -> Optional[str]:
    """Use digits for comparison while preserving the original Excel value separately."""
    original = _text(value)
    return re.sub(r"\D", "", original) if original is not None else None


def review_reason_codes(issues: Iterable[ValidationIssue]) -> Tuple[str, ...]:
    """Keep review causes as stable codes, without raw customer data or messages."""
    reasons = []
    mapping = {
        ("INVALID_PHONE", "phone_original"): "PHONE_INVALID",
        ("REVIEW_REQUIRED", "phone_original"): "PHONE_INVALID",
        ("INVALID_QUANTITY", "quantity"): "QUANTITY_INVALID",
        ("INVALID_DATE", "application_date"): "APPLICATION_DATE_INVALID",
        ("INVALID_DATE", "pickup_date"): "PICKUP_DATE_INVALID",
        ("REVIEW_REQUIRED", "address"): "ADDRESS_CHECK",
        ("REVIEW_REQUIRED", "serial_number"): "SERIAL_CHECK",
    }
    for issue in issues:
        code = mapping.get((issue.code, issue.field))
        if code and code not in reasons:
            reasons.append(code)
    return tuple(reasons)


def preview_recall_applications(
    file_bytes: bytes,
    *,
    source_filename: Optional[str] = None,
    existing_records: Iterable[Mapping[str, Any]] = (),
    upload_batch: Optional[str] = None,
    uploaded_by: Optional[str] = None,
    uploaded_at: Optional[datetime] = None,
) -> ApplicationPreview:
    """Parse the described 11-column application sheet without registering rows.

    Existing records, when provided by a later upload flow, are read only.
    They may contain id, status, serial_number and phone_normalized/original.
    """
    if not file_bytes:
        raise RecallApplicationExcelError("빈 Excel 파일입니다.")
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise RecallApplicationExcelError("읽을 수 없는 .xlsx 파일입니다.") from exc

    try:
        sheet, header_number, indexes, headers = _find_header(workbook)
        if sheet.max_row - header_number > MAX_PREVIEW_ROWS or sheet.max_column > MAX_PREVIEW_COLUMNS:
            raise RecallApplicationExcelError("Excel 미리보기 범위가 너무 큽니다.")
        existing_serials, existing_phones = _index_existing(existing_records)
        seen_serials = {}
        seen_phones = {}
        preview_rows = []
        for raw_row_number, cells in enumerate(
            sheet.iter_rows(min_row=header_number + 1, values_only=True), start=header_number + 1
        ):
            raw_data = {
                _raw_column_key(headers, index): _raw_value(value)
                for index, value in enumerate(cells)
            }
            values = {
                field: cells[index] if index < len(cells) else None
                for field, index in indexes.items()
            }
            row_kind = _classify_row(values, cells)
            if row_kind != "data":
                preview_rows.append(PreviewRow(raw_row_number, row_kind, "excluded", raw_data, {}, ()))
                continue

            data, issues = _validate_values(values, workbook.epoch)
            serial_key = _serial_key(data["serial_number"])
            phone_key = data["phone_normalized"] if _valid_phone(data["phone_normalized"]) else None
            # Serial findings precede phone findings. Neither key is unique by itself.
            if serial_key:
                if serial_key in seen_serials:
                    issues.append(ValidationIssue("DUPLICATE_SERIAL", "serial_number", "파일 내부 시리얼번호 중복", "file", seen_serials[serial_key]))
                if serial_key in existing_serials:
                    issues.extend(_existing_issue("DUPLICATE_SERIAL", "serial_number", record)
                                  for record in existing_serials[serial_key])
                seen_serials.setdefault(serial_key, raw_row_number)
            if phone_key:
                if phone_key in seen_phones:
                    issues.append(ValidationIssue("DUPLICATE_PHONE", "phone_normalized", "파일 내부 연락처 중복", "file", seen_phones[phone_key]))
                if phone_key in existing_phones:
                    issues.extend(_existing_issue("DUPLICATE_PHONE", "phone_normalized", record)
                                  for record in existing_phones[phone_key])
                seen_phones.setdefault(phone_key, raw_row_number)

            codes = {issue.code for issue in issues}
            if codes & {"MISSING_REQUIRED", "INVALID_QUANTITY", "INVALID_DATE"}:
                status = "error"
            elif codes & {"DUPLICATE_SERIAL", "DUPLICATE_PHONE"}:
                status = "duplicate"
            elif codes & {"REVIEW_REQUIRED"}:
                status = "review"
            else:
                status = "valid"
            preview_rows.append(PreviewRow(raw_row_number, row_kind, status, raw_data, data, tuple(issues)))

        return ApplicationPreview(source_filename, sheet.title, header_number,
                                  {field: headers[index] for field, index in indexes.items()},
                                  tuple(preview_rows), upload_batch, uploaded_by, uploaded_at)
    finally:
        workbook.close()


def _find_header(workbook):
    for sheet in workbook.worksheets:
        for row_number, cells in enumerate(sheet.iter_rows(max_row=min(sheet.max_row, 30), values_only=True), start=1):
            headers = [_text(value) for value in cells]
            normalized = [_header_key(value) for value in headers]
            indexes = {}
            for label, field in EXCEL_COLUMNS.items():
                matches = [index for index, value in enumerate(normalized) if value == _header_key(label)]
                if len(matches) > 1:
                    raise RecallApplicationExcelError("중복된 Excel 컬럼입니다: {}".format(label))
                if matches:
                    indexes[field] = matches[0]
            if len(indexes) >= 7 and all(field in indexes for field in REQUIRED_FIELDS):
                missing = [label for label, field in EXCEL_COLUMNS.items() if field not in indexes]
                if missing:
                    raise RecallApplicationExcelError("신청 데이터 Excel 컬럼이 없습니다: {}".format(", ".join(missing)))
                return sheet, row_number, indexes, headers
    raise RecallApplicationExcelError("11개 신청 데이터 컬럼을 가진 시트를 찾지 못했습니다.")


def _header_key(value):
    return "".join((value or "").replace("*", "").split()).casefold()


def _classify_row(values, cells):
    populated = [_text(value) for value in cells if _text(value) is not None]
    if not populated:
        return "blank"
    if len(populated) <= 2 and any(marker in text for text in populated for marker in INSTRUCTION_MARKERS):
        customer_name = _text(values["customer_name"])
        if not customer_name or any(marker in customer_name for marker in INSTRUCTION_MARKERS):
            return "instruction"
    return "data"


def _validate_values(values, epoch):
    data = {field: _text(value) for field, value in values.items()}
    issues = []
    for field in REQUIRED_FIELDS:
        if data[field] is None:
            issues.append(ValidationIssue("MISSING_REQUIRED", field, "필수값이 없습니다."))

    data["phone_normalized"] = normalize_phone(values["phone_original"])
    if data["phone_original"] is not None and not _valid_phone(data["phone_normalized"]):
        issues.append(ValidationIssue("INVALID_PHONE", "phone_original", "연락처를 해석할 수 없어 확인이 필요합니다."))
        issues.append(ValidationIssue("REVIEW_REQUIRED", "phone_original", "연락처 확인이 필요합니다."))

    raw_quantity = values["quantity"]
    data["quantity"] = _positive_integer(raw_quantity)
    if _text(raw_quantity) is not None and data["quantity"] is None:
        issues.append(ValidationIssue("INVALID_QUANTITY", "quantity", "수량은 양의 정수여야 합니다."))

    for field in DATE_FIELDS:
        raw_date = values[field]
        data[field] = _parse_date(raw_date, epoch)
        if _text(raw_date) is not None and data[field] is None:
            issues.append(ValidationIssue("INVALID_DATE", field, "날짜를 해석할 수 없습니다."))

    if data["serial_number"] is None:
        issues.append(ValidationIssue("REVIEW_REQUIRED", "serial_number", "시리얼번호가 없어 확인이 필요합니다."))
    address = data["address"]
    if address and (len(address) < 5 or address.casefold() in {"미정", "없음", "?", "주소불명"}):
        issues.append(ValidationIssue("REVIEW_REQUIRED", "address", "주소 확인이 필요합니다."))
    # Agreement strings remain untouched; canonical statuses are undecided.
    for field in AGREEMENT_FIELDS:
        data[field] = _text(values[field])
    return data, issues


def _text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    result = str(value).strip()
    return result or None


def _raw_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _raw_column_key(headers, index):
    label = headers[index] if index < len(headers) else None
    if not label:
        return "column_{}".format(index + 1)
    return "column_{}:{}".format(index + 1, label) if headers.count(label) > 1 else label


def _valid_phone(value):
    return bool(value and re.fullmatch(r"[0-9]{10,15}", value))


def _positive_integer(value):
    if isinstance(value, bool):
        return None
    text = _text(value)
    if text and re.fullmatch(r"[0-9]+", text) and int(text) > 0:
        return int(text)
    return None


def _parse_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value, epoch=epoch)
            return parsed.date() if isinstance(parsed, datetime) else parsed if isinstance(parsed, date) else None
        except (ValueError, OverflowError):
            return None
    text = _text(value)
    match = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", text or "")
    if match:
        try:
            return date(*(int(part) for part in match.groups()))
        except ValueError:
            return None
    return None


def _serial_key(value):
    return value.strip().casefold() if value else None


def _index_existing(records):
    serials, phones = {}, {}
    for record in records:
        serial = _serial_key(_text(record.get("serial_number")))
        phone = normalize_phone(record.get("phone_normalized") or record.get("phone_original"))
        if serial:
            serials.setdefault(serial, []).append(record)
        if _valid_phone(phone):
            phones.setdefault(phone, []).append(record)
    return serials, phones


def _existing_issue(code, field, record):
    return ValidationIssue(code, field, "기존 데이터와 중복되어 확인이 필요합니다.",
                           source="existing", existing_id=record.get("id"),
                           existing_status=record.get("status"))
