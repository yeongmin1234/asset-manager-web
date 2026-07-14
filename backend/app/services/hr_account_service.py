from datetime import date, datetime, timezone
from io import BytesIO
from typing import Dict, List, Optional, Set, Tuple

from sqlalchemy import or_, select
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.hr_account import HrAccount
from app.schemas.hr_account import (
    HrAccountCreate, HrAccountImportData, HrAccountImportPreviewResponse,
    HrAccountImportPreviewRow, HrAccountImportResponse, HrAccountUpdate,
)


class HrAccountNotFoundError(Exception):
    pass


class HrAccountImportValidationError(Exception):
    pass


HR_IMPORT_COLUMNS = ("department", "name", "dowoffice", "erp", "scm", "nas")
HR_IMPORT_COLUMN_LABELS = {
    "department": "부서", "name": "이름", "dowoffice": "다우오피스",
    "erp": "ERP", "scm": "SCM", "nas": "NAS",
}
HR_IMPORT_ALIASES = {
    "department": ("부서", "부서명", "소속", "소속부서"),
    "name": ("이름", "성명", "직원명", "사용자명"),
    "dowoffice": ("다우오피스", "다우", "다우 ID", "다우ID", "다우 계정", "다우계정"),
    "erp": ("ERP", "ERP ID", "ERPID", "ERP 계정", "ERP계정"),
    "scm": ("SCM", "SCM ID", "SCMID", "SCM 계정", "SCM계정"),
    "nas": ("NAS", "NAS ID", "NASID", "NAS 계정", "NAS계정"),
}
MAX_HR_IMPORT_ROWS = 5000


def list_hr_accounts(db: Session, keyword: Optional[str] = None) -> List[HrAccount]:
    statement = select(HrAccount).where(HrAccount.deleted_at.is_(None))
    normalized = (keyword or "").strip()
    if normalized:
        pattern = "%{}%".format(normalized)
        statement = statement.where(or_(
            HrAccount.department.ilike(pattern), HrAccount.name.ilike(pattern),
            HrAccount.dowoffice.ilike(pattern), HrAccount.erp.ilike(pattern),
            HrAccount.scm.ilike(pattern), HrAccount.nas.ilike(pattern),
        ))
    return list(db.scalars(statement.order_by(HrAccount.created_at.desc(), HrAccount.id.desc())).all())


def get_hr_account(db: Session, account_id: int) -> HrAccount:
    account = db.get(HrAccount, account_id)
    if account is None or account.deleted_at is not None:
        raise HrAccountNotFoundError()
    return account


def create_hr_account(db: Session, payload: HrAccountCreate) -> HrAccount:
    account = HrAccount(**payload.model_dump())
    db.add(account)
    db.commit()
    db.refresh(account)
    return account


def update_hr_account(db: Session, account_id: int, payload: HrAccountUpdate) -> HrAccount:
    account = get_hr_account(db, account_id)
    for key, value in payload.model_dump().items():
        setattr(account, key, value)
    account.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(account)
    return account


def delete_hr_account(db: Session, account_id: int) -> HrAccount:
    account = get_hr_account(db, account_id)
    account.deleted_at = datetime.now(timezone.utc)
    account.updated_at = account.deleted_at
    db.commit()
    db.refresh(account)
    return account


def build_hr_account_import_template() -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "인사업무계정등록양식"
    worksheet.append([HR_IMPORT_COLUMN_LABELS[key] for key in HR_IMPORT_COLUMNS])
    worksheet.append(["총무팀", "홍길동", "hong", "hong", "hong", "hong"])
    fill = PatternFill(fill_type="solid", fgColor="EAF7ED")
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="234D32")
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate((18, 16, 20, 18, 18, 18), start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def preview_hr_accounts_import(db: Session, file_bytes: bytes) -> HrAccountImportPreviewResponse:
    worksheet = _load_hr_import_worksheet(file_bytes)
    matched_indexes, matched_columns = _match_hr_import_headers(worksheet)
    if worksheet.max_row < 2:
        raise HrAccountImportValidationError("등록할 데이터가 없는 빈 엑셀 파일입니다.")
    if worksheet.max_row - 1 > MAX_HR_IMPORT_ROWS:
        raise HrAccountImportValidationError("한 번에 최대 5,000행까지 등록할 수 있습니다.")

    existing_keys = {
        _hr_duplicate_key(account.department, account.name)
        for account in db.scalars(select(HrAccount).where(HrAccount.deleted_at.is_(None))).all()
    }
    seen_keys: Set[Tuple[str, str]] = set()
    rows: List[HrAccountImportPreviewRow] = []
    for row_number in range(2, worksheet.max_row + 1):
        raw_values = [worksheet.cell(row=row_number, column=index + 1).value for index in range(worksheet.max_column)]
        if all(_normalize_hr_import_value(value) is None for value in raw_values):
            continue
        values = {
            key: _normalize_hr_import_value(worksheet.cell(row=row_number, column=index + 1).value)
            if index is not None else None
            for key, index in matched_indexes.items()
        }
        errors = _validate_hr_import_values(values)
        row_status = "error"
        if not errors:
            duplicate_key = _hr_duplicate_key(values["department"], values["name"])
            row_status = "duplicate" if duplicate_key in existing_keys or duplicate_key in seen_keys else "valid"
            seen_keys.add(duplicate_key)
        rows.append(HrAccountImportPreviewRow(
            row_number=row_number,
            status=row_status,
            data=HrAccountImportData(**values),
            errors=errors,
        ))

    if not rows:
        raise HrAccountImportValidationError("등록할 데이터가 없는 빈 엑셀 파일입니다.")
    return HrAccountImportPreviewResponse(
        matched_columns=matched_columns,
        total_count=len(rows),
        valid_count=sum(row.status == "valid" for row in rows),
        duplicate_count=sum(row.status == "duplicate" for row in rows),
        error_count=sum(row.status == "error" for row in rows),
        rows=rows,
    )


def import_hr_accounts(
    db: Session,
    rows: List[HrAccountImportPreviewRow],
    duplicate_policy: str,
) -> HrAccountImportResponse:
    if len(rows) > MAX_HR_IMPORT_ROWS:
        raise HrAccountImportValidationError("한 번에 최대 5,000행까지 등록할 수 있습니다.")
    created_count = updated_count = skipped_count = failed_count = 0
    try:
        existing_accounts = {
            _hr_duplicate_key(account.department, account.name): account
            for account in db.scalars(select(HrAccount).where(HrAccount.deleted_at.is_(None))).all()
        }
        for row in rows:
            values = {key: getattr(row.data, key) for key in HR_IMPORT_COLUMNS}
            errors = _validate_hr_import_values(values)
            if row.status == "error" or errors:
                failed_count += 1
                continue
            key = _hr_duplicate_key(values["department"], values["name"])
            existing = existing_accounts.get(key)
            if existing is not None:
                if duplicate_policy == "skip":
                    skipped_count += 1
                    continue
                for field in ("dowoffice", "erp", "scm", "nas"):
                    setattr(existing, field, values[field])
                existing.updated_at = datetime.now(timezone.utc)
                updated_count += 1
                continue
            account = HrAccount(**values)
            db.add(account)
            existing_accounts[key] = account
            created_count += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    return HrAccountImportResponse(
        created_count=created_count,
        updated_count=updated_count,
        skipped_count=skipped_count,
        failed_count=failed_count,
    )


def _load_hr_import_worksheet(file_bytes: bytes):
    if not file_bytes:
        raise HrAccountImportValidationError("빈 파일은 업로드할 수 없습니다.")
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise HrAccountImportValidationError("올바른 .xlsx 파일이 아니거나 암호화된 파일입니다.") from exc
    if not workbook.sheetnames:
        raise HrAccountImportValidationError("엑셀 파일에 시트가 없습니다.")
    return workbook[workbook.sheetnames[0]]


def _match_hr_import_headers(worksheet) -> Tuple[Dict[str, Optional[int]], Dict[str, str]]:
    raw_headers = [_normalize_hr_import_value(cell.value) for cell in worksheet[1]]
    if not any(raw_headers):
        raise HrAccountImportValidationError("첫 번째 행에 컬럼명이 없습니다.")
    normalized_headers: Dict[str, int] = {}
    duplicate_headers: List[str] = []
    for index, header in enumerate(raw_headers):
        if not header:
            continue
        normalized = _normalize_hr_header(header)
        if normalized in normalized_headers:
            duplicate_headers.append(header)
        else:
            normalized_headers[normalized] = index
    if duplicate_headers:
        raise HrAccountImportValidationError("중복된 엑셀 컬럼명이 있습니다: {}".format(", ".join(duplicate_headers)))

    matched_indexes: Dict[str, Optional[int]] = {}
    matched_columns: Dict[str, str] = {}
    used_indexes: Set[int] = set()
    for key in HR_IMPORT_COLUMNS:
        alias_indexes = {
            normalized_headers[_normalize_hr_header(alias)]
            for alias in HR_IMPORT_ALIASES[key]
            if _normalize_hr_header(alias) in normalized_headers
        }
        if len(alias_indexes) > 1:
            raise HrAccountImportValidationError(
                "{} 컬럼에 해당하는 헤더가 여러 개 있습니다.".format(HR_IMPORT_COLUMN_LABELS[key])
            )
        index = next(iter(alias_indexes), None)
        if index is not None and index in used_indexes:
            raise HrAccountImportValidationError("하나의 엑셀 컬럼이 여러 계정 컬럼에 중복 매칭되었습니다.")
        matched_indexes[key] = index
        if index is not None:
            used_indexes.add(index)
            matched_columns[key] = raw_headers[index]
    missing = [HR_IMPORT_COLUMN_LABELS[key] for key in ("department", "name") if matched_indexes[key] is None]
    if missing:
        raise HrAccountImportValidationError("필수 컬럼이 없습니다: {}".format(", ".join(missing)))
    return matched_indexes, matched_columns


def _normalize_hr_header(value: str) -> str:
    return "".join(str(value).strip().split()).casefold()


def _normalize_hr_import_value(value: object) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        normalized = str(value)
    elif isinstance(value, float) and value.is_integer():
        normalized = str(int(value))
    elif isinstance(value, (date, datetime)):
        normalized = value.isoformat()
    else:
        normalized = str(value)
    normalized = normalized.strip()
    return normalized or None


def _validate_hr_import_values(values: Dict[str, Optional[str]]) -> List[str]:
    errors: List[str] = []
    if not values.get("department"):
        errors.append("부서가 없습니다.")
    if not values.get("name"):
        errors.append("이름이 없습니다.")
    for field in HR_IMPORT_COLUMNS:
        value = values.get(field)
        max_length = 100 if field in ("department", "name") else 200
        if value and len(value) > max_length:
            errors.append("{}은(는) {}자 이하여야 합니다.".format(HR_IMPORT_COLUMN_LABELS[field], max_length))
    return errors


def _hr_duplicate_key(department: Optional[str], name: Optional[str]) -> Tuple[str, str]:
    return ((department or "").strip().casefold(), (name or "").strip().casefold())
