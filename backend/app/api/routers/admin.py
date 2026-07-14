from datetime import date, datetime, timedelta
from io import BytesIO
import logging
from typing import Literal, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy.exc import SQLAlchemyError

from app.core.config import settings
from app.core.auth import require_admin
from app.db.database import get_db
from app.models.user import User
from app.schemas.admin import (
    AdminPasswordRequest,
    AdminPasswordResetRequest,
    AdminPasswordResetResponse,
    AdminPasswordResponse,
    AdminStatusResponse,
    AdminVerifyRequest,
    AdminVerifyResponse,
)
from app.services.activity_log_service import record_activity_log
from app.schemas.login_access_log import LoginAccessLogPage
from app.schemas.audit_log import AuditLogDetail, AuditLogPage
from app.schemas.menu_access_log import MenuAccessLogPage
from app.services.login_access_log_service import get_access_logs
from app.services.audit_log_service import get_audit_log, get_audit_logs, query_audit_logs_for_export, record_audit_log
from app.services.menu_access_log_service import get_menu_access_logs
from app.services.admin_service import (
    AdminPasswordInvalidError,
    AdminPasswordRequiredError,
    AdminPasswordTooShortError,
    AdminResetCodeNotConfiguredError,
    create_admin_auth_token,
    is_admin_password_configured,
    reset_admin_password_with_reset_code,
    set_admin_password,
    verify_admin_reset_code,
    verify_admin_password as verify_stored_admin_password,
)


router = APIRouter(prefix="/admin", tags=["admin"])
logger = logging.getLogger(__name__)


@router.get("/access-logs", response_model=LoginAccessLogPage)
def list_login_access_logs(
    keyword: Optional[str] = Query(default=None),
    username: Optional[str] = Query(default=None),
    event_type: Optional[Literal["login", "logout"]] = Query(default=None),
    login_result: Optional[Literal["success", "failure"]] = Query(default=None),
    access_type: Optional[Literal["internal", "external"]] = Query(default=None),
    start_date: Optional[date] = Query(default=None),
    end_date: Optional[date] = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
) -> LoginAccessLogPage:
    try:
        return get_access_logs(
            db,
            keyword=keyword,
            username=username,
            event_type=event_type,
            login_result=login_result,
            access_type=access_type,
            start_date=start_date,
            end_date=end_date,
            page=page,
            page_size=page_size,
        )
    except SQLAlchemyError as exc:
        logger.exception("Failed to load login access logs")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="접속기록을 불러오는 중 DB 연결에 실패했습니다.",
        ) from exc


@router.get("/menu-access-logs", response_model=MenuAccessLogPage)
def list_menu_access_logs(
    keyword: Optional[str] = Query(default=None),
    username: Optional[str] = Query(default=None),
    menu_key: Optional[str] = Query(default=None),
    access_type: Optional[Literal["internal", "external"]] = Query(default=None),
    start_date: Optional[date] = Query(default=None),
    end_date: Optional[date] = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
) -> MenuAccessLogPage:
    try:
        return get_menu_access_logs(
            db,
            keyword=keyword,
            username=username,
            menu_key=menu_key,
            access_type=access_type,
            start_date=start_date,
            end_date=end_date,
            page=page,
            page_size=page_size,
        )
    except SQLAlchemyError as exc:
        logger.exception("Failed to load menu access logs")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="메뉴 접근 기록을 불러오는 중 DB 연결에 실패했습니다.",
        ) from exc


@router.get("/audit-logs", response_model=AuditLogPage)
def list_audit_logs(
    keyword: Optional[str] = Query(default=None),
    username: Optional[str] = Query(default=None),
    action_type: Optional[Literal["create", "update", "delete", "activate", "deactivate", "permission_change", "excel_import", "export"]] = Query(default=None),
    menu_key: Optional[str] = Query(default=None),
    target_type: Optional[str] = Query(default=None),
    access_type: Optional[Literal["internal", "external"]] = Query(default=None),
    start_date: Optional[date] = Query(default=None),
    end_date: Optional[date] = Query(default=None),
    changed_field: Optional[str] = Query(default=None),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
) -> AuditLogPage:
    try:
        return get_audit_logs(
            db, keyword=keyword, username=username, action_type=action_type,
            menu_key=menu_key, target_type=target_type, access_type=access_type,
            start_date=start_date, end_date=end_date, page=page, page_size=page_size, changed_field=changed_field,
        )
    except SQLAlchemyError as exc:
        logger.exception("Failed to load audit logs")
        raise HTTPException(status_code=503, detail="감사로그를 불러오는 중 DB 연결에 실패했습니다.") from exc


@router.get("/audit-logs/export")
def export_audit_logs(
    request: Request,
    keyword: Optional[str] = Query(default=None), username: Optional[str] = Query(default=None),
    action_type: Optional[str] = Query(default=None), menu_key: Optional[str] = Query(default=None),
    target_type: Optional[str] = Query(default=None), access_type: Optional[str] = Query(default=None),
    start_date: Optional[date] = Query(default=None), end_date: Optional[date] = Query(default=None),
    changed_field: Optional[str] = Query(default=None), db: Session = Depends(get_db),
    current_admin: User = Depends(require_admin),
):
    logs = query_audit_logs_for_export(
        db, keyword=keyword, username=username, action_type=action_type, menu_key=menu_key,
        target_type=target_type, access_type=access_type, start_date=start_date, end_date=end_date,
        changed_field=changed_field, limit=10000,
    )
    if len(logs) > 10000:
        raise HTTPException(status_code=400, detail="다운로드 가능한 최대 건수는 10,000건입니다. 기간 또는 필터를 줄여주세요.")
    output = _build_audit_excel(logs)
    record_audit_log(
        db, request, current_admin, action_type="export", menu_key="audit_logs", menu_name="감사로그",
        target_type="audit_log_export", target_id=None, target_name="감사로그 엑셀 다운로드",
        action_summary="감사로그 {}건을 엑셀로 다운로드했습니다.".format(len(logs)),
    )
    filename = "감사로그_{}_{}.xlsx".format(start_date.strftime("%Y%m%d"), end_date.strftime("%Y%m%d")) if start_date and end_date else "감사로그_{}.xlsx".format(datetime.now().strftime("%Y%m%d_%H%M%S"))
    return StreamingResponse(
        output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''{}".format(quote(filename))},
    )


@router.get("/audit-logs/{audit_log_id}", response_model=AuditLogDetail)
def read_audit_log(audit_log_id: int, db: Session = Depends(get_db)) -> AuditLogDetail:
    item = get_audit_log(db, audit_log_id)
    if item is None:
        raise HTTPException(status_code=404, detail="감사로그를 찾을 수 없습니다.")
    return item


AUDIT_FIELD_LABELS = {
    "department": "부서", "name": "이름", "erp": "ERP", "scm": "SCM", "nas": "NAS",
    "dowoffice": "다우오피스", "menu_permissions": "메뉴 권한", "status": "상태", "role": "역할",
    "is_active": "활성 상태", "title": "제목", "vehicle_number": "차량번호", "owner_name": "사용자",
}
AUDIT_ACTION_LABELS = {"create": "등록", "update": "수정", "delete": "삭제", "activate": "활성화", "deactivate": "비활성화", "permission_change": "권한 변경", "excel_import": "엑셀 일괄등록", "export": "엑셀 다운로드"}
AUDIT_PERMISSION_LABELS = {"dashboard": "대시보드", "assets": "자산 관리", "software": "SW 현황", "company_cars": "법인차량 관리", "hr_list": "인사업무 리스트", "statistics": "통계 / 리포트", "changelog": "변경 이력", "work_manual": "업무설명서", "vendor_contacts": "업체연락처", "expiration_schedules": "점검·만료 관리"}


def _build_audit_excel(logs):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "감사로그"
    headers = ["작업 일시", "사용자 ID", "사용자 이름", "메뉴", "작업 유형", "대상 유형", "대상 이름", "작업 내용", "변경 항목", "변경 전", "변경 후", "접속 IP", "접속 구분", "브라우저", "운영체제"]
    sheet.append(headers)
    for log in logs:
        fields = log.changed_fields or []
        sheet.append([
            log.occurred_at.strftime("%Y-%m-%d %H:%M:%S") if log.occurred_at else "", log.username,
            log.user_name, log.menu_name, AUDIT_ACTION_LABELS.get(log.action_type, log.action_type), log.target_type, log.target_name or "",
            log.action_summary, "\n".join(_field_label(field) for field in fields),
            _detail_text(log.before_data, fields), _detail_text(log.after_data, fields),
            log.ip_address or "", "내부망" if log.access_type == "internal" else "외부망",
            log.browser or "", log.operating_system or "",
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="E5E7EB")
    widths = [20, 15, 15, 20, 16, 20, 24, 40, 24, 45, 45, 18, 12, 14, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO(); workbook.save(stream); stream.seek(0); return stream


def _field_label(field):
    return AUDIT_FIELD_LABELS.get(field, field)


def _detail_text(data, fields):
    if not data:
        return ""
    keys = fields or list(data.keys())
    return "\n".join("{}: {}".format(_field_label(key), _display_value(data.get(key), key)) for key in keys if key in data)


def _display_value(value, field=None):
    if value is None:
        return "-"
    if value == "":
        return "(빈 값)"
    if isinstance(value, list):
        if field == "menu_permissions":
            return ", ".join(AUDIT_PERMISSION_LABELS.get(str(item), str(item)) for item in value)
        return ", ".join(str(item) for item in value)
    return str(value)


@router.get("/status", response_model=AdminStatusResponse)
def get_admin_status(db: Session = Depends(get_db)) -> AdminStatusResponse:
    return AdminStatusResponse(configured=is_admin_password_configured(db))


@router.post("/password", response_model=AdminPasswordResponse)
def update_admin_password(
    payload: AdminPasswordRequest,
    db: Session = Depends(get_db),
) -> AdminPasswordResponse:
    try:
        set_admin_password(
            db,
            payload.new_password,
            current_password=payload.current_password or None,
        )
    except AdminPasswordTooShortError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="관리자 비밀번호는 6자 이상이어야 합니다.",
        ) from exc
    except AdminPasswordRequiredError as exc:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="현재 비밀번호를 입력해주세요.",
        ) from exc
    except AdminPasswordInvalidError as exc:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="현재 비밀번호가 올바르지 않습니다.",
        ) from exc

    return AdminPasswordResponse(ok=True, configured=True)


@router.post("/password/reset", response_model=AdminPasswordResetResponse)
def reset_admin_password(
    payload: AdminPasswordResetRequest,
    request: Request,
    db: Session = Depends(get_db),
) -> AdminPasswordResetResponse:
    try:
        if payload.new_password != payload.confirm_password:
            raise AdminPasswordInvalidError()
        if not verify_admin_reset_code(payload.reset_code):
            raise AdminPasswordInvalidError()
        reset_admin_password_with_reset_code(db, payload.new_password)
    except AdminResetCodeNotConfiguredError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="관리자 비밀번호 초기화 코드가 설정되지 않았습니다.",
        ) from exc
    except (AdminPasswordTooShortError, AdminPasswordInvalidError) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="초기화 코드 또는 입력값을 확인해 주세요.",
        ) from exc

    record_activity_log(
        db,
        menu_name="설정",
        action_type="admin-password-reset",
        target_type="admin",
        target_id=None,
        target_name="관리자 비밀번호",
        actor_ip=request.client.host if request.client else None,
        user_agent=request.headers.get("user-agent"),
        summary="관리자 비밀번호 웹 초기화",
        before_data=None,
        after_data=None,
    )
    db.commit()
    return AdminPasswordResetResponse(
        ok=True,
        message="관리자 비밀번호가 초기화되었습니다. 새 비밀번호로 다시 인증해 주세요.",
    )


@router.post("/verify", response_model=AdminVerifyResponse)
def verify_admin_password(
    payload: AdminVerifyRequest,
    db: Session = Depends(get_db),
) -> AdminVerifyResponse:
    if not is_admin_password_configured(db):
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="관리자 비밀번호가 설정되지 않았습니다.",
        )

    if not verify_stored_admin_password(db, payload.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="비밀번호가 올바르지 않습니다.",
        )

    expires_at = datetime.utcnow() + timedelta(minutes=_get_auth_minutes())
    return AdminVerifyResponse(
        ok=True,
        token=create_admin_auth_token(db, expires_at),
        expires_at=expires_at,
    )


def _get_auth_minutes() -> int:
    try:
        return max(1, int(settings.admin_auth_minutes))
    except (TypeError, ValueError):
        return 60
